from __future__ import annotations

import re
from pathlib import Path
from zipfile import ZipFile

from docx import Document


DOC = Path(r"D:\Arch_Projects\HYP_天津_武清区六街 L-9 地块_Gary\01_Office\01 采购文件附件八：【设计任务书】规划建筑方案设计任务书.docx")
TERMS = re.compile(
    r"容积率|建筑密度|绿地率|建筑高度|限高|建筑面积|用地面积|总图|户型|套型|停车|车位|层数|"
    r"洋房|叠拼|改善|刚需|住宅|配套|地块|红线|退界|退线|日照|消防|人车分流|地下|公共服务|"
    r"幼儿园|商业|学校|河道|道路|指标|设计要求|规划条件"
)


def clean(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def main() -> None:
    doc = Document(str(DOC))
    print("paragraphs", len(doc.paragraphs), "tables", len(doc.tables), "sections", len(doc.sections))
    with ZipFile(DOC) as archive:
        media = [name for name in archive.namelist() if name.startswith("word/media/")]
        embedded = [name for name in archive.namelist() if name.startswith("word/embeddings/")]
    print("media", media)
    print("embedded", embedded)
    print("\n## 匹配段落")
    for i, paragraph in enumerate(doc.paragraphs):
        text = clean(paragraph.text)
        if text and TERMS.search(text):
            print(f"P{i}: {text}")
    print("\n## 表格")
    for ti, table in enumerate(doc.tables):
        print(f"TABLE {ti} rows={len(table.rows)} cols={len(table.columns)}")
        for ri, row in enumerate(table.rows):
            values = [clean(cell.text) for cell in row.cells]
            line = " | ".join(values)
            if line and (TERMS.search(line) or ri < 3):
                print(f"T{ti}R{ri}: {line}")

    print("\n## 关键段落原文区间")
    for start, end in ((40, 50), (72, 83), (90, 106), (108, 125)):
        print(f"RANGE {start}-{end}")
        for i in range(start, min(end + 1, len(doc.paragraphs))):
            text = clean(doc.paragraphs[i].text)
            if text:
                print(f"P{i}: {text}")


if __name__ == "__main__":
    main()
